"""Unit tests for app.core.extractor — extract_text() and ExtractionError.

Fixtures are generated programmatically using PyMuPDF and python-docx so
the tests never depend on committed binary files.  All tests are pure unit
tests: no database, no HTTP client, no disk I/O.

Coverage:
    PDF — valid single-page, valid multi-page, no-text (image-only),
          encrypted, corrupt bytes
    DOCX — valid single paragraph, valid multi-paragraph, empty document,
           corrupt bytes
    Common — unsupported extension, ExtractionError is a ValueError subclass
"""

import io

import fitz  # PyMuPDF
import pytest
from docx import Document

from app.core.extractor import ExtractionError, extract_text

# ---------------------------------------------------------------------------
# Fixture factories
# (called at module level so the bytes are built once per test session)
# ---------------------------------------------------------------------------


def _make_pdf(text: str) -> bytes:
    """Create a minimal single-page text PDF containing *text*."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _make_multipage_pdf(pages: list[str]) -> bytes:
    """Create a multi-page PDF with one text block per page."""
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _make_image_only_pdf() -> bytes:
    """Create a PDF with one blank page (no text layer — simulates scanned page)."""
    doc = fitz.open()
    doc.new_page()  # blank page, no text
    data = doc.tobytes()
    doc.close()
    return data


def _make_encrypted_pdf() -> bytes:
    """Create a password-protected PDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "secret")
    data = doc.tobytes(
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="owner",
        user_pw="user",
    )
    doc.close()
    return data


def _make_docx(paragraphs: list[str]) -> bytes:
    """Create a DOCX with the given paragraph strings."""
    buf = io.BytesIO()
    d = Document()
    for para in paragraphs:
        d.add_paragraph(para)
    d.save(buf)
    return buf.getvalue()


def _make_empty_docx() -> bytes:
    """Create a DOCX with no text content (default empty document)."""
    buf = io.BytesIO()
    Document().save(buf)
    return buf.getvalue()


# Pre-build the fixtures once
PDF_VALID = _make_pdf("John Doe\nSoftware Engineer\nPython, FastAPI")
PDF_MULTIPAGE = _make_multipage_pdf(["Page one content", "Page two content"])
PDF_IMAGE_ONLY = _make_image_only_pdf()
PDF_ENCRYPTED = _make_encrypted_pdf()
PDF_CORRUPT = b"%PDF-1.4 this is not a real pdf \x00\xff\xfe"
DOCX_VALID = _make_docx(["Jane Smith", "Data Scientist", "Skills: Python, ML"])
DOCX_MULTIPARA = _make_docx(["Paragraph one.", "Paragraph two.", "Paragraph three."])
DOCX_EMPTY = _make_empty_docx()
DOCX_CORRUPT = b"PK\x03\x04 this is not a valid zip/docx \x00\xff\xfe"


# ===========================================================================
# ExtractionError type
# ===========================================================================


class TestExtractionError:
    def test_is_value_error_subclass(self):
        err = ExtractionError("something went wrong")
        assert isinstance(err, ValueError)

    def test_message_preserved(self):
        err = ExtractionError("test message")
        assert str(err) == "test message"


# ===========================================================================
# PDF extraction
# ===========================================================================


class TestExtractPdf:
    def test_valid_pdf_returns_text(self):
        result = extract_text(PDF_VALID, "pdf")
        assert isinstance(result, str)
        assert len(result) > 0

    def test_valid_pdf_contains_expected_content(self):
        result = extract_text(PDF_VALID, "pdf")
        assert "John Doe" in result

    def test_valid_pdf_result_is_not_just_whitespace(self):
        result = extract_text(PDF_VALID, "pdf")
        assert result.strip() != ""

    def test_multipage_pdf_contains_all_pages(self):
        result = extract_text(PDF_MULTIPAGE, "pdf")
        assert "Page one content" in result
        assert "Page two content" in result

    def test_multipage_pdf_pages_separated(self):
        result = extract_text(PDF_MULTIPAGE, "pdf")
        # Pages joined by double newline
        assert "\n\n" in result

    def test_image_only_pdf_raises_extraction_error(self):
        with pytest.raises(ExtractionError) as exc_info:
            extract_text(PDF_IMAGE_ONLY, "pdf")
        assert "scanned" in str(exc_info.value).lower() or "no text" in str(exc_info.value).lower()

    def test_encrypted_pdf_raises_extraction_error(self):
        with pytest.raises(ExtractionError) as exc_info:
            extract_text(PDF_ENCRYPTED, "pdf")
        assert "password" in str(exc_info.value).lower() or "protected" in str(exc_info.value).lower()

    def test_corrupt_pdf_raises_extraction_error(self):
        with pytest.raises(ExtractionError):
            extract_text(PDF_CORRUPT, "pdf")

    def test_pdf_extension_uppercase_works(self):
        """Extension matching should be case-insensitive."""
        result = extract_text(PDF_VALID, "PDF")
        assert "John Doe" in result


# ===========================================================================
# DOCX extraction
# ===========================================================================


class TestExtractDocx:
    def test_valid_docx_returns_text(self):
        result = extract_text(DOCX_VALID, "docx")
        assert isinstance(result, str)
        assert len(result) > 0

    def test_valid_docx_contains_expected_content(self):
        result = extract_text(DOCX_VALID, "docx")
        assert "Jane Smith" in result

    def test_multiparagraph_docx_contains_all_content(self):
        result = extract_text(DOCX_MULTIPARA, "docx")
        assert "Paragraph one." in result
        assert "Paragraph two." in result
        assert "Paragraph three." in result

    def test_multiparagraph_docx_paragraphs_separated_by_newline(self):
        result = extract_text(DOCX_MULTIPARA, "docx")
        assert "\n" in result

    def test_empty_docx_raises_extraction_error(self):
        with pytest.raises(ExtractionError) as exc_info:
            extract_text(DOCX_EMPTY, "docx")
        assert "empty" in str(exc_info.value).lower()

    def test_corrupt_docx_raises_extraction_error(self):
        with pytest.raises(ExtractionError):
            extract_text(DOCX_CORRUPT, "docx")

    def test_docx_extension_uppercase_works(self):
        result = extract_text(DOCX_VALID, "DOCX")
        assert "Jane Smith" in result


# ===========================================================================
# Unsupported extensions
# ===========================================================================


class TestUnsupportedExtension:
    def test_txt_raises_extraction_error(self):
        with pytest.raises(ExtractionError) as exc_info:
            extract_text(b"some text content", "txt")
        assert "unsupported" in str(exc_info.value).lower()

    def test_png_raises_extraction_error(self):
        with pytest.raises(ExtractionError):
            extract_text(b"\x89PNG\r\n", "png")

    def test_empty_ext_raises_extraction_error(self):
        with pytest.raises(ExtractionError):
            extract_text(b"data", "")

    def test_error_message_mentions_supported_formats(self):
        with pytest.raises(ExtractionError) as exc_info:
            extract_text(b"data", "xyz")
        detail = str(exc_info.value).lower()
        assert "pdf" in detail or "docx" in detail
